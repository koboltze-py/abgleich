"""
Abgleich-Workflow: vergleicht Ursprungsplanung und Tagesdienstpläne je
Mitarbeiter und ermöglicht das geordnete Abarbeiten der gefundenen
Abweichungen (neue Dienste, geänderte Zeiten, entfallene Dienste). Zeigt
außerdem den vollen Monat je Mitarbeiter editierbar an sowie einen
Kalender-Tab mit dem daraus resultierenden, exportierbaren Monatsplan.
"""
import calendar
import os
from datetime import date

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QComboBox, QPushButton,
    QListWidget, QListWidgetItem, QTableWidget, QTableWidgetItem, QFrame,
    QAbstractItemView, QSplitter, QMessageBox, QInputDialog, QTabWidget,
    QFileDialog, QLineEdit,
)
from PySide6.QtCore import Qt, QUrl, QMarginsF
from PySide6.QtGui import QColor, QDesktopServices, QFont, QTextDocument, QPageSize
from PySide6.QtPrintSupport import QPrinter

from config import (
    FIORI_TEXT, FIORI_SUCCESS, FIORI_WARNING, FIORI_ERROR, FIORI_BORDER,
    QUELLE_STAERKEMELDUNG,
)
from gui.styles import table_style, button_primary, button_secondary, card_style
from gui.widgets import ZeitBereichEditor
from gui.monatsuebersicht import MONATSNAMEN, WOCHENTAGE_KURZ
from gui.dienstplan_ansicht import DienstplanAnsichtWidget
from functions.abgleich_service import (
    berechne_abgleich, set_eintrag_status, set_status_fuer_mitarbeiter,
    get_letzte_position, set_letzte_position, get_verfuegbare_monate,
    get_manuelle_zuordnungen, manuell_zusammenfuehren, manuelle_zuordnung_aufheben,
    manuell_trennen, monatsplan_fuer_mitarbeiter, setze_manuellen_dienst,
    voller_monatsplan,
    ART_HINZUGEFUEGT, ART_ZEIT_GEAENDERT, ART_ENTFALLEN, ART_UNVERAENDERT, ART_KEINE_DATEN,
    STATUS_OFFEN, STATUS_ERLEDIGT,
)
from functions.dienste_service import get_dokumente

ART_LABEL = {
    ART_HINZUGEFUEGT: "Hinzugekommen",
    ART_ZEIT_GEAENDERT: "Zeit geändert",
    ART_ENTFALLEN: "Entfallen",
    ART_UNVERAENDERT: "Unverändert",
    ART_KEINE_DATEN: "–",
}
ART_FARBE = {
    ART_HINZUGEFUEGT: FIORI_SUCCESS,
    ART_ZEIT_GEAENDERT: FIORI_WARNING,
    ART_ENTFALLEN: FIORI_ERROR,
}
_DIFF_ARTEN = {ART_HINZUGEFUEGT, ART_ZEIT_GEAENDERT, ART_ENTFALLEN}


def _hex_zu_rgba(hex_farbe: str, alpha: int) -> QColor:
    c = QColor(hex_farbe)
    c.setAlpha(alpha)
    return c


def _hex_zu_css_rgba(hex_farbe: str, alpha: int) -> str:
    c = QColor(hex_farbe)
    return f"rgba({c.red()}, {c.green()}, {c.blue()}, {alpha})"


class AbgleichWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._ergebnis = None
        self._aktueller_mitarbeiter_id: int | None = None
        self._kalender_plan: list = []
        self._kalender_jahr: int | None = None
        self._kalender_monat: int | None = None
        self._aufbauen()
        self.aktualisieren()

    # ------------------------------------------------------------------
    def _aufbauen(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(12)

        titel = QLabel("Abgleich – Ursprungsplanung vs. Tagesdienstpläne")
        titel.setStyleSheet(f"font-size: 18px; font-weight: bold; color: {FIORI_TEXT};")
        layout.addWidget(titel)

        hinweis = QLabel(
            "Zeigt je Mitarbeiter, welche Dienste gegenüber der Ursprungsplanung hinzugekommen, "
            "zeitlich verändert oder entfallen sind. Bearbeitete Änderungen können hier als "
            "erledigt markiert werden, sobald sie in die externe Liste übertragen wurden. Im Tab "
            "\"Kalender\" steht der daraus resultierende Monatsplan, exportierbar als Excel oder PDF."
        )
        hinweis.setWordWrap(True)
        hinweis.setStyleSheet(f"color: {FIORI_TEXT};")
        layout.addWidget(hinweis)

        filter_zeile = QHBoxLayout()
        filter_zeile.addWidget(QLabel("Monat:"))
        self._monat_combo = QComboBox()
        self._monat_combo.currentIndexChanged.connect(self._monat_gewechselt)
        filter_zeile.addWidget(self._monat_combo)

        filter_zeile.addSpacing(20)
        filter_zeile.addWidget(QLabel("Word-Dokument:"))
        self._dokument_combo = QComboBox()
        self._dokument_combo.setMinimumWidth(260)
        filter_zeile.addWidget(self._dokument_combo)
        btn_dokument_oeffnen = QPushButton("Öffnen")
        btn_dokument_oeffnen.setStyleSheet(button_secondary())
        btn_dokument_oeffnen.clicked.connect(self._dokument_oeffnen)
        filter_zeile.addWidget(btn_dokument_oeffnen)

        filter_zeile.addStretch()
        self._status_label = QLabel("")
        self._status_label.setStyleSheet(f"color: {FIORI_TEXT};")
        filter_zeile.addWidget(self._status_label)
        layout.addLayout(filter_zeile)

        self._haupt_tabs = QTabWidget()
        layout.addWidget(self._haupt_tabs, stretch=1)
        self._haupt_tabs.addTab(self._bearbeitung_tab_erstellen(), "Abgleich bearbeiten")
        self._haupt_tabs.addTab(self._kalender_tab_erstellen(), "Kalender")

    def _bearbeitung_tab_erstellen(self) -> QWidget:
        tab = QWidget()
        tab_layout = QVBoxLayout(tab)
        tab_layout.setContentsMargins(0, 0, 0, 0)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        tab_layout.addWidget(splitter, stretch=1)

        # --- linke Seite: Mitarbeiterliste mit Fortschritt ---
        links = QFrame()
        links.setStyleSheet(card_style())
        links_layout = QVBoxLayout(links)
        links_layout.setContentsMargins(10, 10, 10, 10)
        links_kopf = QLabel("Mitarbeiter mit Abweichungen")
        links_kopf.setStyleSheet(f"font-weight: bold; color: {FIORI_TEXT}; border: none;")
        links_layout.addWidget(links_kopf)
        self._mitarbeiter_suche = QLineEdit()
        self._mitarbeiter_suche.setPlaceholderText("Name suchen…")
        self._mitarbeiter_suche.textChanged.connect(self._mitarbeiter_liste_filtern)
        links_layout.addWidget(self._mitarbeiter_suche)
        self._mitarbeiter_liste = QListWidget()
        self._mitarbeiter_liste.setStyleSheet(f"border: 1px solid {FIORI_BORDER}; border-radius: 4px;")
        self._mitarbeiter_liste.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self._mitarbeiter_liste.currentItemChanged.connect(self._mitarbeiter_gewaehlt)
        links_layout.addWidget(self._mitarbeiter_liste, stretch=1)
        nav_zeile = QHBoxLayout()
        btn_naechster_offen = QPushButton("Nächster mit offenen Punkten")
        btn_naechster_offen.setStyleSheet(button_secondary())
        btn_naechster_offen.clicked.connect(self._naechster_offener)
        nav_zeile.addWidget(btn_naechster_offen)
        links_layout.addLayout(nav_zeile)

        merge_zeile = QHBoxLayout()
        btn_manuell_mergen = QPushButton("Ausgewählte zusammenführen (selbe Person)")
        btn_manuell_mergen.setStyleSheet(button_secondary())
        btn_manuell_mergen.clicked.connect(self._manuell_zusammenfuehren)
        merge_zeile.addWidget(btn_manuell_mergen)
        btn_manuell_trennen = QPushButton("Manuelle Verknüpfung aufheben")
        btn_manuell_trennen.setStyleSheet(button_secondary())
        btn_manuell_trennen.clicked.connect(self._manuelle_verknuepfung_aufheben)
        merge_zeile.addWidget(btn_manuell_trennen)
        links_layout.addLayout(merge_zeile)

        trenn_zeile = QHBoxLayout()
        btn_namen_trennen = QPushButton("Namen aus Zusammenführung trennen…")
        btn_namen_trennen.setStyleSheet(button_secondary())
        btn_namen_trennen.clicked.connect(self._namen_trennen)
        trenn_zeile.addWidget(btn_namen_trennen)
        links_layout.addLayout(trenn_zeile)
        splitter.addWidget(links)

        # --- rechte Seite: voller Monat des gewählten Mitarbeiters (editierbar) ---
        rechts = QFrame()
        rechts.setStyleSheet(card_style())
        rechts_layout = QVBoxLayout(rechts)
        rechts_layout.setContentsMargins(10, 10, 10, 10)

        self._mitarbeiter_titel = QLabel("Bitte Mitarbeiter auswählen")
        self._mitarbeiter_titel.setStyleSheet(f"font-size: 15px; font-weight: bold; color: {FIORI_TEXT}; border: none;")
        rechts_layout.addWidget(self._mitarbeiter_titel)

        self._varianten_label = QLabel("")
        self._varianten_label.setWordWrap(True)
        self._varianten_label.setStyleSheet(f"color: {FIORI_TEXT}; border: none;")
        self._varianten_label.setVisible(False)
        rechts_layout.addWidget(self._varianten_label)

        self._warnung_label = QLabel("")
        self._warnung_label.setWordWrap(True)
        self._warnung_label.setStyleSheet(f"color: {FIORI_ERROR}; font-weight: bold; border: none;")
        self._warnung_label.setVisible(False)
        rechts_layout.addWidget(self._warnung_label)

        hinweis = QLabel(
            "Zeigt den vollen Monat, auch Tage ohne jede Meldung. In der Spalte \"Tatsächlich\" kann "
            "über die Kontrollbox eine Zeit aktiviert und per Uhrzeit-Auswahl eingetragen oder "
            "korrigiert werden (Kontrollbox abwählen zum Löschen)."
        )
        hinweis.setWordWrap(True)
        hinweis.setStyleSheet(f"color: {FIORI_TEXT}; border: none;")
        rechts_layout.addWidget(hinweis)

        self._tabelle = QTableWidget(0, 6)
        self._tabelle.setHorizontalHeaderLabels(
            ["Datum", "Wochentag", "Art", "Ursprünglich", "Tatsächlich", "Erledigt"]
        )
        self._tabelle.setStyleSheet(table_style())
        self._tabelle.setEditTriggers(
            QAbstractItemView.EditTrigger.DoubleClicked | QAbstractItemView.EditTrigger.EditKeyPressed
        )
        self._tabelle.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self._tabelle.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self._tabelle.horizontalHeader().setStretchLastSection(False)
        self._tabelle.itemChanged.connect(self._zelle_geaendert)
        rechts_layout.addWidget(self._tabelle, stretch=1)

        aktion_zeile = QHBoxLayout()
        btn_ausgewaehlt_erledigt = QPushButton("Ausgewählte als erledigt markieren")
        btn_ausgewaehlt_erledigt.setStyleSheet(button_secondary())
        btn_ausgewaehlt_erledigt.clicked.connect(lambda: self._markieren(STATUS_ERLEDIGT, nur_auswahl=True))
        aktion_zeile.addWidget(btn_ausgewaehlt_erledigt)

        btn_alle_erledigt = QPushButton("Mitarbeiter komplett erledigt")
        btn_alle_erledigt.setStyleSheet(button_primary())
        btn_alle_erledigt.clicked.connect(lambda: self._markieren(STATUS_ERLEDIGT, nur_auswahl=False))
        aktion_zeile.addWidget(btn_alle_erledigt)
        aktion_zeile.addStretch()
        rechts_layout.addLayout(aktion_zeile)

        splitter.addWidget(rechts)

        # --- ganz rechts: Tagesdienstplan-Ansicht (nur zur Kontrolle, kein Export) ---
        dienstplan_rahmen = QFrame()
        dienstplan_rahmen.setStyleSheet(card_style())
        dienstplan_rahmen_layout = QVBoxLayout(dienstplan_rahmen)
        dienstplan_rahmen_layout.setContentsMargins(10, 10, 10, 10)
        dienstplan_rahmen_layout.addWidget(DienstplanAnsichtWidget())
        splitter.addWidget(dienstplan_rahmen)

        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 3)
        splitter.setStretchFactor(2, 2)
        return tab

    def _kalender_tab_erstellen(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(10)

        hinweis = QLabel(
            "Finaler Monatsplan aus dem Abgleich: tatsächliche Zeiten (inkl. von Hand nachgetragener "
            "Werte in \"Abgleich bearbeiten\"), sonst die Ursprungsplanung, wo nichts anderes bekannt ist."
        )
        hinweis.setWordWrap(True)
        hinweis.setStyleSheet(f"color: {FIORI_TEXT};")
        layout.addWidget(hinweis)

        self._kalender_tabelle = QTableWidget()
        self._kalender_tabelle.setStyleSheet(table_style())
        self._kalender_tabelle.setAlternatingRowColors(True)
        self._kalender_tabelle.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._kalender_tabelle.verticalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignLeft)
        layout.addWidget(self._kalender_tabelle, stretch=1)

        export_zeile = QHBoxLayout()
        btn_excel = QPushButton("Als Excel exportieren…")
        btn_excel.setStyleSheet(button_secondary())
        btn_excel.clicked.connect(self._als_excel_exportieren)
        export_zeile.addWidget(btn_excel)
        btn_pdf = QPushButton("Als PDF exportieren…")
        btn_pdf.setStyleSheet(button_secondary())
        btn_pdf.clicked.connect(self._als_pdf_exportieren)
        export_zeile.addWidget(btn_pdf)
        export_zeile.addStretch()
        layout.addLayout(export_zeile)

        return tab

    # ------------------------------------------------------------------
    def _monate_laden(self):
        self._monat_combo.blockSignals(True)
        self._monat_combo.clear()
        monate = get_verfuegbare_monate()
        for jahr, monat in monate:
            self._monat_combo.addItem(f"{MONATSNAMEN[monat - 1]} {jahr}", (jahr, monat))
        self._monat_combo.blockSignals(False)

    def _monat_gewechselt(self):
        self._daten_laden()

    def aktualisieren(self):
        """Neu berechnen, z. B. nach einem Import."""
        aktuelle_auswahl = self._monat_combo.currentData()
        self._monate_laden()
        if aktuelle_auswahl:
            idx = self._monat_combo.findData(aktuelle_auswahl)
            if idx >= 0:
                self._monat_combo.setCurrentIndex(idx)
        self._daten_laden()

    # ------------------------------------------------------------------
    def _daten_laden(self):
        auswahl = self._monat_combo.currentData()
        if not auswahl:
            self._ergebnis = None
            self._mitarbeiter_liste.clear()
            self._tabelle.setRowCount(0)
            self._dokument_combo.clear()
            self._kalender_plan = []
            self._kalender_tabelle.setRowCount(0)
            self._kalender_tabelle.setColumnCount(0)
            self._mitarbeiter_titel.setText("Kein Monat mit Ursprungsplanung UND Tagesdienstplänen vorhanden.")
            self._status_label.setText("")
            return

        jahr, monat = auswahl
        self._ergebnis = berechne_abgleich(jahr, monat)
        self._dokumente_laden(jahr, monat)
        self._kalender_befuellen(jahr, monat)

        self._mitarbeiter_liste.blockSignals(True)
        self._mitarbeiter_liste.clear()
        for ma in self._ergebnis.mitarbeiter:
            self._mitarbeiter_liste.addItem(self._mitarbeiter_eintrag(ma))
        self._mitarbeiter_liste.blockSignals(False)
        self._mitarbeiter_liste_filtern(self._mitarbeiter_suche.text())

        gesamt_offen = sum(ma.anzahl_offen for ma in self._ergebnis.mitarbeiter)
        self._status_label.setText(
            f"{len(self._ergebnis.mitarbeiter)} Mitarbeiter mit Abweichungen, {gesamt_offen} offen"
        )

        letzte_id = get_letzte_position(jahr, monat)
        ziel_index = 0
        if letzte_id is not None:
            for i, ma in enumerate(self._ergebnis.mitarbeiter):
                if ma.mitarbeiter_id == letzte_id:
                    ziel_index = i
                    break
        if self._mitarbeiter_liste.count():
            self._mitarbeiter_liste.setCurrentRow(ziel_index)
        else:
            self._mitarbeiter_titel.setText("Keine Abweichungen gefunden – Ursprungsplanung und "
                                             "Tagesdienstpläne stimmen für diesen Monat überein.")
            self._warnung_label.setVisible(False)
            self._tabelle.setRowCount(0)

    def _mitarbeiter_eintrag(self, ma) -> QListWidgetItem:
        praefix = "⚠ " if ma.kein_gegenstueck else ""
        text = f"{praefix}{ma.name}  ({ma.anzahl_offen}/{ma.anzahl_gesamt} offen)"
        item = QListWidgetItem(text)
        item.setData(Qt.ItemDataRole.UserRole, ma.mitarbeiter_id)
        if ma.anzahl_offen == 0:
            item.setForeground(QColor(FIORI_SUCCESS))
        elif ma.kein_gegenstueck:
            item.setForeground(QColor(FIORI_ERROR))
        return item

    # ------------------------------------------------------------------
    def _mitarbeiter_gewaehlt(self, aktuell: QListWidgetItem, _vorher: QListWidgetItem):
        if not aktuell or not self._ergebnis:
            return
        mid = aktuell.data(Qt.ItemDataRole.UserRole)
        self._aktueller_mitarbeiter_id = mid
        ma = next((m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id == mid), None)
        if not ma:
            return

        self._mitarbeiter_titel.setText(ma.name)
        if len(ma.varianten) > 1:
            self._varianten_label.setText("Zusammengeführt aus: " + ", ".join(ma.varianten))
            self._varianten_label.setVisible(True)
        else:
            self._varianten_label.setVisible(False)

        if ma.nur_ursprungsplanung:
            self._warnung_label.setText(
                "⚠ Für diesen Mitarbeiter liegen in diesem Monat nur Ursprungsplanung-Daten vor – "
                "in den Tagesdienstplänen wurde kein Eintrag gefunden. Möglicherweise stimmt der "
                "Name zwischen beiden Quellen nicht überein."
            )
            self._warnung_label.setVisible(True)
        elif ma.nur_staerkemeldung:
            self._warnung_label.setText(
                "⚠ Für diesen Mitarbeiter liegen in diesem Monat nur Tagesdienstplan-Daten vor – "
                "in der Ursprungsplanung wurde kein Eintrag gefunden. Möglicherweise stimmt der "
                "Name zwischen beiden Quellen nicht überein."
            )
            self._warnung_label.setVisible(True)
        else:
            self._warnung_label.setVisible(False)

        self._tabelle_befuellen(ma)

        if self._ergebnis:
            set_letzte_position(self._ergebnis.jahr, self._ergebnis.monat, mid)

    def _tabelle_befuellen(self, ma):
        jahr, monat = self._ergebnis.jahr, self._ergebnis.monat
        plan = monatsplan_fuer_mitarbeiter(jahr, monat, ma.mitarbeiter_ids)
        self._tabelle.blockSignals(True)
        self._tabelle.setRowCount(len(plan))
        for zeile, tag in enumerate(plan):
            wochentag = WOCHENTAGE_KURZ[date.fromisoformat(tag.datum).weekday()]
            ursprung = f"{tag.ursprung_start} – {tag.ursprung_end}" if tag.ursprung_start else "–"

            werte = [tag.datum, wochentag, ART_LABEL[tag.art], ursprung]
            for spalte, wert in enumerate(werte):
                item = QTableWidgetItem(wert)
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                farbe = ART_FARBE.get(tag.art)
                if farbe:
                    item.setBackground(_hex_zu_rgba(farbe, 30))
                self._tabelle.setItem(zeile, spalte, item)

            editor = ZeitBereichEditor(
                tag.tatsaechlich_start, tag.tatsaechlich_end,
                lambda zeit, datum=tag.datum: self._tatsaechlich_zeit_geaendert(datum, zeit),
            )
            farbe = ART_FARBE.get(tag.art)
            if farbe:
                css = _hex_zu_css_rgba(farbe, 60)
                editor.setStyleSheet(f"background-color: {css};")
            self._tabelle.setCellWidget(zeile, 4, editor)

            status_item = QTableWidgetItem()
            if tag.art in _DIFF_ARTEN:
                status_item.setFlags(
                    (status_item.flags() | Qt.ItemFlag.ItemIsUserCheckable) & ~Qt.ItemFlag.ItemIsEditable
                )
                status_item.setCheckState(
                    Qt.CheckState.Checked if tag.status == STATUS_ERLEDIGT else Qt.CheckState.Unchecked
                )
                status_item.setData(Qt.ItemDataRole.UserRole, (tag.datum, tag.art))
            else:
                status_item.setFlags(Qt.ItemFlag.NoItemFlags)
            self._tabelle.setItem(zeile, 5, status_item)

        self._tabelle.resizeColumnsToContents()
        self._tabelle.blockSignals(False)

    # ------------------------------------------------------------------
    def _zelle_geaendert(self, item: QTableWidgetItem):
        if self._aktueller_mitarbeiter_id is None:
            return
        if item.column() == 5:
            self._status_geaendert(item)

    def _status_geaendert(self, item: QTableWidgetItem):
        daten = item.data(Qt.ItemDataRole.UserRole)
        if not daten:
            return
        datum, art = daten
        status = STATUS_ERLEDIGT if item.checkState() == Qt.CheckState.Checked else STATUS_OFFEN
        set_eintrag_status(self._aktueller_mitarbeiter_id, datum, art, status)
        self._eintrag_status_aktualisieren(datum, art, status)
        self._mitarbeiter_liste_aktualisieren()

    def _tatsaechlich_zeit_geaendert(self, datum: str, zeit: tuple[str, str] | None):
        if self._aktueller_mitarbeiter_id is None:
            return
        setze_manuellen_dienst(self._aktueller_mitarbeiter_id, datum, zeit)
        self.aktualisieren()

    def _eintrag_status_aktualisieren(self, datum: str, art: str, status: str):
        if not self._ergebnis or self._aktueller_mitarbeiter_id is None:
            return
        ma = next((m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id == self._aktueller_mitarbeiter_id), None)
        if not ma:
            return
        for e in ma.eintraege:
            if e.datum == datum and e.art == art:
                e.status = status

    def _mitarbeiter_liste_aktualisieren(self):
        aktuelle_zeile = self._mitarbeiter_liste.currentRow()
        self._mitarbeiter_liste.blockSignals(True)
        for i, ma in enumerate(self._ergebnis.mitarbeiter):
            self._mitarbeiter_liste.takeItem(i)
            self._mitarbeiter_liste.insertItem(i, self._mitarbeiter_eintrag(ma))
        self._mitarbeiter_liste.setCurrentRow(aktuelle_zeile)
        self._mitarbeiter_liste.blockSignals(False)
        self._mitarbeiter_liste_filtern(self._mitarbeiter_suche.text())
        gesamt_offen = sum(ma.anzahl_offen for ma in self._ergebnis.mitarbeiter)
        self._status_label.setText(
            f"{len(self._ergebnis.mitarbeiter)} Mitarbeiter mit Abweichungen, {gesamt_offen} offen"
        )

    # ------------------------------------------------------------------
    def _markieren(self, status: str, nur_auswahl: bool):
        if not self._ergebnis or self._aktueller_mitarbeiter_id is None:
            return
        ma = next((m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id == self._aktueller_mitarbeiter_id), None)
        if not ma:
            return

        if nur_auswahl:
            # Tabelle zeigt jetzt den vollen Monat - über (Datum, Art) der
            # markierten Zeilen (Spalte "Erledigt") die passenden Einträge finden
            zeilen = sorted({i.row() for i in self._tabelle.selectedIndexes()})
            schluessel = set()
            for z in zeilen:
                daten = self._tabelle.item(z, 5).data(Qt.ItemDataRole.UserRole)
                if daten:
                    schluessel.add(daten)
            eintraege = [e for e in ma.eintraege if (e.datum, e.art) in schluessel]
        else:
            eintraege = ma.eintraege

        if not eintraege:
            return
        set_status_fuer_mitarbeiter(ma.mitarbeiter_id, eintraege, status)
        for e in eintraege:
            e.status = status
        self._tabelle_befuellen(ma)
        self._mitarbeiter_liste_aktualisieren()

    def _naechster_offener(self):
        if not self._ergebnis or not self._ergebnis.mitarbeiter:
            return
        anzahl = len(self._ergebnis.mitarbeiter)
        start = self._mitarbeiter_liste.currentRow()
        for schritt in range(1, anzahl + 1):
            index = (start + schritt) % anzahl
            if self._ergebnis.mitarbeiter[index].anzahl_offen > 0:
                self._mitarbeiter_liste.setCurrentRow(index)
                return

    def _mitarbeiter_liste_filtern(self, text: str):
        text = text.strip().lower()
        for i in range(self._mitarbeiter_liste.count()):
            item = self._mitarbeiter_liste.item(i)
            item.setHidden(bool(text) and text not in item.text().lower())

    # ------------------------------------------------------------------
    def _ausgewaehlte_mitarbeiter(self) -> list:
        if not self._ergebnis:
            return []
        ids = [item.data(Qt.ItemDataRole.UserRole) for item in self._mitarbeiter_liste.selectedItems()]
        return [m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id in ids]

    def _manuell_zusammenfuehren(self):
        ausgewaehlt = self._ausgewaehlte_mitarbeiter()
        if len(ausgewaehlt) < 2:
            QMessageBox.information(
                self, "Zusammenführen",
                "Bitte mindestens zwei Mitarbeiter in der Liste auswählen (Strg+Klick), "
                "die dieselbe Person sind."
            )
            return
        basis = ausgewaehlt[0].mitarbeiter_id
        for weitere in ausgewaehlt[1:]:
            manuell_zusammenfuehren(basis, weitere.mitarbeiter_id)
        self.aktualisieren()

    def _manuelle_verknuepfung_aufheben(self):
        if not self._ergebnis or self._aktueller_mitarbeiter_id is None:
            return
        ma = next((m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id == self._aktueller_mitarbeiter_id), None)
        if not ma or len(ma.mitarbeiter_ids) < 2:
            QMessageBox.information(self, "Verknüpfung aufheben", "Dieser Mitarbeiter hat keine manuelle Verknüpfung.")
            return
        eigene_ids = set(ma.mitarbeiter_ids)
        betroffen = [
            (a, b) for a, b in get_manuelle_zuordnungen()
            if a in eigene_ids and b in eigene_ids
        ]
        if not betroffen:
            QMessageBox.information(
                self, "Verknüpfung aufheben",
                "Diese Zusammenführung wurde automatisch erkannt und nicht manuell erstellt - "
                "sie kann hier nicht aufgehoben werden."
            )
            return
        for a, b in betroffen:
            manuelle_zuordnung_aufheben(a, b)
        self.aktualisieren()

    def _namen_trennen(self):
        if not self._ergebnis or self._aktueller_mitarbeiter_id is None:
            return
        ma = next((m for m in self._ergebnis.mitarbeiter if m.mitarbeiter_id == self._aktueller_mitarbeiter_id), None)
        if not ma or len(ma.mitglieder) < 2:
            QMessageBox.information(
                self, "Namen trennen",
                "Dieser Mitarbeiter besteht nur aus einem Namen - es gibt nichts zu trennen."
            )
            return
        namen = [name for name, _mid in ma.mitglieder]
        auswahl, ok = QInputDialog.getItem(
            self, "Namen trennen",
            "Welcher Name gehört NICHT zu den anderen und soll wieder als eigener "
            "Mitarbeiter behandelt werden?",
            namen, 0, False,
        )
        if not ok or not auswahl:
            return
        heraustrennen_id = next(mid for name, mid in ma.mitglieder if name == auswahl)
        for name, mid in ma.mitglieder:
            if mid != heraustrennen_id:
                manuell_trennen(heraustrennen_id, mid)
        self.aktualisieren()

    # ------------------------------------------------------------------
    def _dokumente_laden(self, jahr: int, monat: int):
        """Füllt das Dropdown mit den für den Monat importierten Word-Dokumenten."""
        self._dokument_combo.clear()
        dokumente = get_dokumente(quelle=QUELLE_STAERKEMELDUNG, jahr=jahr, monat=monat)
        if not dokumente:
            self._dokument_combo.addItem("Keine Word-Dokumente für diesen Monat", None)
            return
        for dok in dokumente:
            self._dokument_combo.addItem(dok["dateiname"], dok["dateipfad"])

    def _dokument_oeffnen(self):
        pfad = self._dokument_combo.currentData()
        if not pfad:
            return
        if not os.path.isfile(pfad):
            QMessageBox.warning(self, "Datei nicht gefunden", f"Die Datei wurde nicht gefunden:\n{pfad}")
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(pfad))

    # ------------------------------------------------------------------
    @staticmethod
    def _finale_zeit(tag) -> tuple[str, str] | None:
        if tag.tatsaechlich_start:
            return tag.tatsaechlich_start, tag.tatsaechlich_end
        if tag.ursprung_start:
            return tag.ursprung_start, tag.ursprung_end
        return None

    def _kalender_befuellen(self, jahr: int, monat: int):
        self._kalender_jahr, self._kalender_monat = jahr, monat
        self._kalender_plan = voller_monatsplan(jahr, monat)

        anzahl_tage = calendar.monthrange(jahr, monat)[1]
        tabelle = self._kalender_tabelle
        tabelle.clear()
        tabelle.setRowCount(len(self._kalender_plan))
        tabelle.setColumnCount(anzahl_tage)
        tabelle.setVerticalHeaderLabels([m.name for m in self._kalender_plan])

        header_labels = []
        for tag in range(1, anzahl_tage + 1):
            wochentag = WOCHENTAGE_KURZ[date(jahr, monat, tag).weekday()]
            header_labels.append(f"{tag:02d}\n{wochentag}")
        tabelle.setHorizontalHeaderLabels(header_labels)

        for zeile, mitarbeiter in enumerate(self._kalender_plan):
            for spalte, tag in enumerate(mitarbeiter.tage):
                zeit = self._finale_zeit(tag)
                if not zeit:
                    continue
                item = QTableWidgetItem(f"{zeit[0]}-{zeit[1]}")
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                farbe = ART_FARBE.get(tag.art)
                if farbe:
                    item.setBackground(_hex_zu_rgba(farbe, 35))
                tabelle.setItem(zeile, spalte, item)

        tabelle.resizeColumnsToContents()
        tabelle.horizontalHeader().setMinimumSectionSize(58)

    def _als_excel_exportieren(self):
        if not self._kalender_plan:
            QMessageBox.information(self, "Export", "Kein Monatsplan zum Exportieren vorhanden.")
            return
        vorschlag = f"Monatsplan_{self._kalender_jahr:04d}-{self._kalender_monat:02d}.xlsx"
        pfad, _ = QFileDialog.getSaveFileName(self, "Monatsplan als Excel speichern", vorschlag, "Excel-Dateien (*.xlsx)")
        if not pfad:
            return
        try:
            self._excel_schreiben(pfad)
        except Exception as exc:
            QMessageBox.warning(self, "Export fehlgeschlagen", str(exc))
            return
        QMessageBox.information(self, "Export erfolgreich", f"Monatsplan gespeichert:\n{pfad}")

    def _excel_schreiben(self, pfad: str):
        from openpyxl import Workbook

        jahr, monat = self._kalender_jahr, self._kalender_monat
        anzahl_tage = calendar.monthrange(jahr, monat)[1]

        wb = Workbook()
        ws = wb.active
        ws.title = "Monatsplan"
        ws.cell(row=1, column=1, value="Mitarbeiter")
        for tag in range(1, anzahl_tage + 1):
            wochentag = WOCHENTAGE_KURZ[date(jahr, monat, tag).weekday()]
            ws.cell(row=1, column=tag + 1, value=f"{tag:02d}.{wochentag}")

        for zeile, mitarbeiter in enumerate(self._kalender_plan, start=2):
            ws.cell(row=zeile, column=1, value=mitarbeiter.name)
            for spalte, tag in enumerate(mitarbeiter.tage, start=2):
                zeit = self._finale_zeit(tag)
                if zeit:
                    ws.cell(row=zeile, column=spalte, value=f"{zeit[0]}-{zeit[1]}")

        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 22
        wb.save(pfad)

    def _als_pdf_exportieren(self):
        if not self._kalender_plan:
            QMessageBox.information(self, "Export", "Kein Monatsplan zum Exportieren vorhanden.")
            return
        vorschlag = f"Monatsplan_{self._kalender_jahr:04d}-{self._kalender_monat:02d}.pdf"
        pfad, _ = QFileDialog.getSaveFileName(self, "Monatsplan als PDF speichern", vorschlag, "PDF-Dateien (*.pdf)")
        if not pfad:
            return
        try:
            self._pdf_schreiben(pfad)
        except Exception as exc:
            QMessageBox.warning(self, "Export fehlgeschlagen", str(exc))
            return
        QMessageBox.information(self, "Export erfolgreich", f"Monatsplan gespeichert:\n{pfad}")

    def _pdf_schreiben(self, pfad: str):
        jahr, monat = self._kalender_jahr, self._kalender_monat
        anzahl_tage = calendar.monthrange(jahr, monat)[1]
        zellen_stil = "white-space:nowrap; border:1px solid #999; padding:2px 6px;"
        FONT_PT = 8

        drucker = QPrinter(QPrinter.PrinterMode.HighResolution)
        drucker.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        drucker.setOutputFileName(pfad)
        layout = drucker.pageLayout()
        layout.setPageSize(QPageSize(QPageSize.PageSizeId.A3))
        layout.setOrientation(layout.Orientation.Landscape)
        layout.setMargins(QMarginsF(8, 8, 8, 8))
        drucker.setPageLayout(layout)
        seiten_breite = drucker.pageRect(QPrinter.Unit.Point).width()

        def tabelle_html(start: int, ende: int) -> str:
            kopf = f"<tr><th style='{zellen_stil}'>Mitarbeiter</th>" + "".join(
                f"<th style='{zellen_stil}'>{tag:02d}<br/>{WOCHENTAGE_KURZ[date(jahr, monat, tag).weekday()].upper()}</th>"
                for tag in range(start, ende + 1)
            ) + "</tr>"
            zeilen_html = []
            for mitarbeiter in self._kalender_plan:
                zellen = [f"<td style='{zellen_stil}'>{mitarbeiter.name}</td>"]
                for tag in mitarbeiter.tage[start - 1:ende]:
                    zeit = self._finale_zeit(tag)
                    text = f"{zeit[0]}\u2013{zeit[1]}" if zeit else ""
                    zellen.append(f"<td style='{zellen_stil}'>{text}</td>")
                zeilen_html.append("<tr>" + "".join(zellen) + "</tr>")
            return (
                f"<table cellspacing='0' cellpadding='0' style='border-collapse:collapse; font-size:{FONT_PT}pt;'>"
                f"{kopf}{''.join(zeilen_html)}</table>"
            )

        def passt_in_seite(html_tabelle: str) -> bool:
            test_dok = QTextDocument()
            test_dok.setDefaultFont(QFont("Segoe UI", FONT_PT))
            test_dok.setHtml(f"<html><body>{html_tabelle}</body></html>")
            return test_dok.idealWidth() <= seiten_breite

        # Ein volles Monatsraster (~30 Tage) ist selbst auf A3 querformat zu
        # breit, um die Uhrzeiten ohne Zeilenumbruch lesbar darzustellen -
        # daher in Tages-Blöcke aufteilen, die je auf eine eigene Seite passen.
        # Die Blockgröße wird anhand der tatsächlichen Textbreite ermittelt,
        # damit auf jedem Rechner (andere Schriftmetriken) nichts abgeschnitten wird.
        bloecke = []
        start = 1
        while start <= anzahl_tage:
            ende = anzahl_tage
            html_tabelle = tabelle_html(start, ende)
            while ende > start and not passt_in_seite(html_tabelle):
                ende -= 1
                html_tabelle = tabelle_html(start, ende)

            seitenumbruch = "page-break-before:always;" if bloecke else ""
            bloecke.append(
                f"<div style='{seitenumbruch}'>"
                f"<h3>Monatsplan {MONATSNAMEN[monat - 1]} {jahr} – Tage {start}–{ende}</h3>"
                f"{html_tabelle}</div>"
            )
            start = ende + 1

        html = "<html><body>" + "".join(bloecke) + "</body></html>"

        dokument = QTextDocument()
        dokument.setDefaultFont(QFont("Segoe UI", FONT_PT))
        dokument.setHtml(html)
        dokument.print_(drucker)


