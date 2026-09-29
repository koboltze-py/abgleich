"""
Parser für Ursprungsplanung-Excel-Dateien (CareMan-Dienstplan-Export, .xlsx).

Format: eine Zeile pro Mitarbeiter und Diensttag mit Datum, Name
("Nachname, Vorname"), Dienst-Code (z. B. "T", "N10", "DN3"), Beginn-/
Ende-Zeitpunkt und "Ebene" (Dienstplan/Abweichung/Arbeitszeit). Titel-,
Gruppen- ("[Stamm FH]") und Fußzeilen werden anhand fehlender/unpassender
Werte übersprungen, die Spaltenlage der Datenspalten ist unabhängig von
den (teils über mehrere Spalten verschmolzenen) Kopfzeilen fix.
"""
from datetime import datetime

from openpyxl import load_workbook

from functions.staerkemeldung_parser import DienstEintrag, ParseErgebnis

# feste Datenspalten des CareMan-Exports (0-basiert): Datum, Name, Dienst, Beginn, Ende, Ebene
_SPALTE_DATUM = 2
_SPALTE_NAME = 3
_SPALTE_DIENST = 4
_SPALTE_BEGINN = 7
_SPALTE_ENDE = 8
_SPALTE_EBENE = 9


def _name_normalisieren(name: str) -> str:
    """"Nachname, Vorname" -> "Nachname V." (Format wie in den Tagesdienstplänen,
    wo überwiegend nur der Nachname bzw. Nachname + Vorname-Kürzel steht)."""
    name = name.strip()
    if "," in name:
        nachname, vorname = name.split(",", 1)
        nachname, vorname = nachname.strip(), vorname.strip()
        if vorname:
            return f"{nachname} {vorname[0].upper()}."
        return nachname
    return name


def parse_datei(dateipfad: str) -> ParseErgebnis:
    """Liest eine Ursprungsplanung-.xlsx-Datei und liefert alle Dienst-Einträge."""
    try:
        wb = load_workbook(dateipfad, data_only=True, read_only=True)
    except Exception as exc:
        return ParseErgebnis(dateipfad, None, None, [], fehler=f"Datei konnte nicht geöffnet werden: {exc}")

    ws = wb.active
    eintraege: list[DienstEintrag] = []
    von_datum = None
    bis_datum = None

    for row in ws.iter_rows():
        if len(row) <= _SPALTE_EBENE:
            continue
        datum_wert = row[_SPALTE_DATUM].value
        name_wert = row[_SPALTE_NAME].value
        beginn_wert = row[_SPALTE_BEGINN].value
        ende_wert = row[_SPALTE_ENDE].value

        if not isinstance(datum_wert, datetime) or not isinstance(name_wert, str) or not name_wert.strip():
            continue
        if not isinstance(beginn_wert, datetime) or not isinstance(ende_wert, datetime):
            continue

        name = _name_normalisieren(name_wert)
        dienst_code = str(row[_SPALTE_DIENST].value or "Sonstige").strip()
        ebene = row[_SPALTE_EBENE].value
        ebene = str(ebene).strip() if ebene else None

        start_datum = beginn_wert.date()
        end_datum = ende_wert.date()
        dauer = int((ende_wert - beginn_wert).total_seconds() // 60)

        eintraege.append(DienstEintrag(
            kategorie=dienst_code, name=name,
            start_zeit=beginn_wert.strftime("%H:%M"), end_zeit=ende_wert.strftime("%H:%M"),
            datum=start_datum, end_datum=end_datum, dauer_minuten=dauer, ebene=ebene,
        ))

        if von_datum is None or start_datum < von_datum:
            von_datum = start_datum
        if bis_datum is None or start_datum > bis_datum:
            bis_datum = start_datum

    wb.close()

    if not eintraege:
        return ParseErgebnis(dateipfad, None, None, [], fehler="Keine Dienst-Einträge gefunden (Format unbekannt)")

    return ParseErgebnis(dateipfad, von_datum, bis_datum, eintraege)
